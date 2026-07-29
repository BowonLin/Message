import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

file_path = 'letter.xlsx'
df = pd.read_excel(file_path, header=None)

# 定義目標名單
name_list = [
    "蔡佳妏", "蔡竣丞", "蔡智森", "陳柏宇", "陳佑瑀", "陳吟昕", 
    "歐書佑", "郭品言", "郭名軒", "游欣穎", "吳洛昕", "林柏宇", 
    "吳芯苡", "林宥安", "盧潔儀", "呂硏瑉", "陳玉倩", "朱品潔", 
    "王梓行", "吳奎諒", "黃子薇", "黃士紘", "黃冠瑜", "張凱傑", 
    "張詠程", "童偀淇", "朱怡臻", "吳欣恩"
]

# 資料結構初始化
recipient_data = {name: [] for name in name_list}

for index, row in df.iterrows():
    sender = row.iloc[0] 
    if pd.isna(sender) or str(sender).strip() == '':
        sender = "匿名"  # 容錯：若寫信人沒名字，自動填入匿名
    else:
        sender = str(sender).strip()
    
    # 過濾掉所有空白格，只留下有實體字內容的
    cols_with_vals = [(df.columns[c], row.iloc[c]) for c in range(1, len(row)) if pd.notna(row.iloc[c]) and str(row.iloc[c]).strip() != ""]
    
    # 掃描有字的所有格子
    for idx, (col_name, val) in enumerate(cols_with_vals):
        val_str = str(val).strip()
        
        # 當格子內容在我們的人名清單內
        if val_str in name_list:
            possible_contents = []
            
            # 因為已經壓縮空白，所以真正的內文必定是緊鄰在它的右邊或左邊
            if idx + 1 < len(cols_with_vals): # 右邊相鄰有字的格子
                possible_contents.append((idx + 1, cols_with_vals[idx + 1][1], cols_with_vals[idx + 1][0]))
            if idx - 1 >= 0: # 左邊相鄰有字的格子
                possible_contents.append((idx - 1, cols_with_vals[idx - 1][1], cols_with_vals[idx - 1][0]))
            
            # 確留下來的「不是人名」而是信件內文
            valid_contents = [c for c in possible_contents if str(c[1]).strip() not in name_list]
            
            message = ""
            if valid_contents:
                if len(valid_contents) == 1:
                    message = valid_contents[0][1]
                else:
                    # 如果左右都有內容，優先匹配對齊原本欄位型態的（to_whom 找 say_it）
                    matched = [c for c in valid_contents if ('to_whom' in col_name and 'say_it' in c[2]) or ('say_it' in col_name and 'to_whom' in c[2])]
                    if matched:
                        message = matched[0][1]
                    else:
                        # 否則選字數最多的那一格作為信件內容
                        message = max(valid_contents, key=lambda x: len(str(x[1])))[1]
            
            if message:
                recipient_data[val_str].append((sender, str(message).strip()))

# 建立 Excel 活頁簿與視覺設計設定
wb = openpyxl.Workbook()
default_sheet = wb.active
wb.remove(default_sheet)

font_title = Font(name='Microsoft JhengHei', size=14, bold=True, color='FFFFFF')
font_message = Font(name='Microsoft JhengHei', size=11, bold=False, color='333333')
font_by = Font(name='Microsoft JhengHei', size=10, italic=True, color='666666')
fill_title = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')
thin_border = Border(
    left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
)

# 寫入資料到各工作表 (Sheet)
for recipient in name_list:
    letters = recipient_data[recipient]
    if not letters:
        continue
        
    ws = wb.create_sheet(title=recipient)
    ws.views.sheetView[0].showGridLines = True
    ws.column_dimensions['A'].width = 70
    
    # 頂部主題大標題
    ws.cell(row=1, column=1, value=f"寫給 {recipient} 的信")
    cell_title = ws['A1']
    cell_title.font = font_title
    cell_title.fill = fill_title
    cell_title.alignment = align_center
    ws.row_dimensions[1].height = 40
    
    current_row = 2
    for sender, message in letters:
        # 寫入內文
        ws.cell(row=current_row, column=1, value=message)
        ws.cell(row=current_row, column=1).font = font_message
        ws.cell(row=current_row, column=1).alignment = align_left
        ws.cell(row=current_row, column=1).border = thin_border
        
        # 依據字數動態給予高度
        num_lines = max(len(str(message).split('\n')), (len(str(message)) // 50) + 1)
        ws.row_dimensions[current_row].height = max(24, num_lines * 18)
        current_row += 1
        
        # 寫入 By 署名
        ws.cell(row=current_row, column=1, value=f"By {sender}")
        ws.cell(row=current_row, column=1).font = font_by
        ws.cell(row=current_row, column=1).alignment = align_right
        ws.cell(row=current_row, column=1).border = thin_border
        ws.row_dimensions[current_row].height = 20
        current_row += 1
        
        # 格子間空一行
        ws.cell(row=current_row, column=1, value="")
        ws.row_dimensions[current_row].height = 15
        current_row += 1

# 儲存檔案
output_file = '實力.xlsx'
wb.save(output_file)
print(f"全部抓齊！已成功輸出檔案：{output_file}")
